import json
import os
import sys

from gatherer import gather_dataframes
from loaders import load_dataframe

DATA_EXTENSIONS = {".csv", ".xlsx", ".xls", ".json"}

def discover_data_files(directory):
    """Returns a sorted list of data-file names (CSV/Excel/JSON) directly
    inside `directory`. Not recursive, and skips hidden files/folders -
    used when "/pandex gather" is run with no filenames, to auto-combine
    every data file sitting in the current project."""
    found = []
    for entry in os.scandir(directory):
        if not entry.is_file():
            continue
        if entry.name.startswith("."):
            continue
        if os.path.splitext(entry.name)[1].lower() in DATA_EXTENSIONS:
            found.append(entry.name)
    return sorted(found)

def resolve_input_paths(argv, cwd):
    """Given sys.argv[1:] and the current working directory, returns the
    list of files gather should combine, or raises ValueError with a
    message meant to be shown to the user as-is.

    - 2+ filenames given: use them, same as always.
    - Exactly 1 filename given: still an error, gather needs at least 2.
    - No filenames given: auto-discover every data file in the project
      folder and combine those instead.
    """
    if len(argv) >= 2:
        return list(argv)

    if len(argv) == 1:
        raise ValueError(
            "gather needs at least 2 files to combine. Either list them "
            "(python gather.py file1.csv file2.csv), or run it with no "
            "filenames to auto-combine every data file in this folder."
        )

    discovered = discover_data_files(cwd)
    if len(discovered) < 2:
        raise ValueError(
            f"No filenames were given, so gather looked for data files in "
            f"this folder and found {len(discovered)}. It needs at least "
            f"2 to combine. Add more data files here, or list files "
            f"explicitly: python gather.py file1.csv file2.csv"
        )
    return discovered

def write_last_gathered_state(gathered_file_path, cwd):
    """Records the most recent successful gather output so that "/pandex
    clean" run with no filename can automatically pick up where gather
    left off. Best-effort: if this can't be written for any reason, gather
    still succeeds, it just won't be chainable."""
    if not gathered_file_path:
        return
    try:
        state_dir = os.path.join(cwd, ".pandex")
        os.makedirs(state_dir, exist_ok=True)
        with open(os.path.join(state_dir, "state.json"), "w", encoding="utf-8") as f:
            json.dump({"last_gathered": gathered_file_path}, f)
    except OSError:
        pass

def build_output_folder(first_path):
    """<first file>_gathered_results, next to the first input file (not in the
    current working directory)."""
    first_base = os.path.splitext(first_path)[0]
    folder = f"{first_base}_gathered_results"
    os.makedirs(folder, exist_ok=True)
    return folder

def add_excel_autofilter(worksheet, row_count, column_count):
    from openpyxl.utils import get_column_letter
    if row_count == 0 or column_count == 0:
        return
    last_column_letter = get_column_letter(column_count)
    last_row = row_count + 1
    worksheet.auto_filter.ref = f"A1:{last_column_letter}{last_row}"
    for col_idx, col_name in enumerate(worksheet[1], start=1):
        max_length = max(
            [len(str(col_name.value))] +
            [len(str(cell.value)) for cell in worksheet[get_column_letter(col_idx)][1:200] if cell.value is not None]
        )
        worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 60)

def gather_files(paths):
    dataframes = []
    filenames = []
    file_notes = {}
    for path in paths:
        df, _dup_cols, notes = load_dataframe(path)
        dataframes.append(df)
        filenames.append(os.path.basename(path))
        if notes:
            file_notes[os.path.basename(path)] = notes

    mode, result, report = gather_dataframes(dataframes, filenames)
    if file_notes:
        report["file_notes"] = file_notes

    output_folder = build_output_folder(paths[0])
    output_path = os.path.join(output_folder, "gathered.xlsx")

    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)

    if mode == "joined":
        sheet = workbook.create_sheet("gathered")
        sheet.append(list(result.columns))
        for row in result.itertuples(index=False):
            sheet.append(list(row))
        add_excel_autofilter(sheet, len(result), len(result.columns))
    else:
        for sheet_name, df in result.items():
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(list(df.columns))
            for row in df.itertuples(index=False):
                sheet.append(list(row))
            add_excel_autofilter(sheet, len(df), len(df.columns))

    workbook.save(output_path)

    report["file"] = ", ".join(paths)
    report["output_folder"] = output_folder
    report["gathered_file"] = output_path
    return report

def friendly_error_message(e, paths):
    if isinstance(e, FileNotFoundError):
        return f"Couldn't find one of these files: {', '.join(paths)}. Check the file names and make sure they're all in this folder."
    if isinstance(e, PermissionError):
        return "Couldn't open one of these files - it might be open in another program (like Excel). Close it and try again."
    message = str(e)
    return message

if __name__ == "__main__":
    try:
        file_paths = resolve_input_paths(sys.argv[1:], os.getcwd())
    except ValueError as e:
        print(json.dumps({"error": str(e)}))
        sys.exit(1)

    try:
        result = gather_files(file_paths)
        if len(sys.argv) == 1:
            result["auto_discovered_files"] = file_paths
        print(json.dumps(result, indent=2))
        write_last_gathered_state(result.get("gathered_file"), os.getcwd())
    except Exception as e:
        print(json.dumps({"error": friendly_error_message(e, file_paths)}))
        sys.exit(1)
