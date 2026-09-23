import json
import os
import sys

import pandas as pd
from cleaner import clean_dataframe
from loaders import load_dataframe


def build_output_folder(path):
    base, _ext = os.path.splitext(path)
    folder = f"{base}_cleaned_results"
    os.makedirs(folder, exist_ok=True)
    return folder

def add_excel_autofilter(path, row_count, column_count):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    worksheet = workbook.active

    last_column_letter = get_column_letter(column_count)
    last_row = row_count + 1
    worksheet.auto_filter.ref = f"A1:{last_column_letter}{last_row}"

    for col_idx, col_name in enumerate(worksheet[1], start=1):
        max_length = max(
            [len(str(col_name.value))] +
            [len(str(cell.value)) for cell in worksheet[get_column_letter(col_idx)][1:200] if cell.value is not None]
        )
        worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 60)

    workbook.save(path)

def save_dataframe(df, path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        df.to_csv(path, index=False)
    elif ext in (".xlsx", ".xls"):
        df.to_excel(path, index=False)
        add_excel_autofilter(path, len(df), len(df.columns))
    elif ext == ".json":
        df.to_json(path, orient="records", indent=2)
    else:
        raise ValueError(f"Unsupported file type '{ext}' for saving.")

def clean_file(path):
    df, duplicate_column_names, notes = load_dataframe(path)

    if len(df) == 0:
        return {
            "file": path,
            "warning": "This file has columns but zero data rows - nothing to clean."
        }

    _base, ext = os.path.splitext(path)
    cleaned_df, report, duplicates_df, missing_df, empty_rows_df, conflicts_df = clean_dataframe(df)

    output_folder = build_output_folder(path)

    # The cleaned file, and all review files, are always saved as .xlsx with
    # AutoFilter dropdowns on every column and auto-sized widths - so the
    # user can filter/sort by region, date, gender, status, etc. directly in
    # Excel, regardless of what format the original input file was.
    cleaned_path = os.path.join(output_folder, "cleaned.xlsx")
    save_dataframe(cleaned_df, cleaned_path)
    report["file"] = path
    report["output_folder"] = output_folder
    report["cleaned_file"] = cleaned_path

    if len(duplicates_df) > 0:
        duplicates_path = os.path.join(output_folder, "duplicates.xlsx")
        save_dataframe(duplicates_df, duplicates_path)
        report["duplicates_file"] = duplicates_path

    if len(conflicts_df) > 0:
        conflicts_path = os.path.join(output_folder, "conflicts.xlsx")
        save_dataframe(conflicts_df, conflicts_path)
        report["conflicts_file"] = conflicts_path

    if len(missing_df) > 0:
        missing_path = os.path.join(output_folder, "missing_values.xlsx")
        save_dataframe(missing_df, missing_path)
        report["missing_values_file"] = missing_path

    if len(empty_rows_df) > 0:
        empty_path = os.path.join(output_folder, "empty_rows.xlsx")
        save_dataframe(empty_rows_df, empty_path)
        report["empty_rows_file"] = empty_path

    if duplicate_column_names:
        report["duplicate_column_names_found"] = duplicate_column_names
    report.update(notes)

    return report

def friendly_error_message(e, path):
    if isinstance(e, FileNotFoundError):
        return f"Couldn't find a file at '{path}'. Check the file name and make sure it's in this folder."

    if isinstance(e, PermissionError):
        return f"Couldn't open '{path}' - it might be open in another program (like Excel). Close it and try again."

    if isinstance(e, pd.errors.EmptyDataError):
        return "This file has no columns or rows - there is no data to clean."

    if isinstance(e, pd.errors.ParserError):
        return f"'{path}' doesn't look like a valid, well-formed file. It may be corrupted or use an unusual format."

    if isinstance(e, UnicodeDecodeError):
        return f"Couldn't read the text in '{path}' - it may use an unusual character encoding."

    if isinstance(e, IsADirectoryError):
        return f"'{path}' is a folder, not a file. Point this at a specific file instead."

    message = str(e)
    lower_message = message.lower()
    if "zip file" in lower_message or "not a zip file" in lower_message or "file format cannot be determined" in lower_message or "engine manually" in lower_message:
        return f"'{path}' doesn't look like a valid Excel file. It may be corrupted or actually a different file type with a .xlsx extension."
    if "expecting value" in lower_message or "json" in lower_message:
        return f"'{path}' doesn't look like valid JSON. Check the file's formatting."

    return message

def read_last_gathered_state(cwd):
    """Best-effort read of the file gather last wrote to (see
    write_last_gathered_state in gather.py). Returns None if there's no
    record yet, or it can't be read for any reason."""
    state_path = os.path.join(cwd, ".pandex", "state.json")
    try:
        with open(state_path, "r", encoding="utf-8") as f:
            return json.load(f).get("last_gathered")
    except (OSError, json.JSONDecodeError):
        return None

def resolve_input_path(argv, cwd):
    """Given sys.argv[1:] and the current working directory, returns the
    file clean should operate on, or raises ValueError with a message
    meant to be shown to the user as-is.

    - One filename given: use it.
    - No filename given: fall back to whatever gather last combined, so
      "/pandex clean" alone can pick up straight where "/pandex gather"
      left off.
    - Anything else (2+ filenames): clean only ever works on one file at a
      time, so this is a usage error.
    """
    if len(argv) == 1:
        return argv[0]

    if len(argv) > 1:
        raise ValueError("Usage: python clean.py <path-to-file>")

    last_gathered = read_last_gathered_state(cwd)
    if not last_gathered:
        raise ValueError(
            "No filename was given, and there's no record of a previous "
            "gather to fall back on. Run 'python clean.py <path-to-file>', "
            "or run gather first."
        )

    candidate = last_gathered if os.path.isabs(last_gathered) else os.path.join(cwd, last_gathered)
    candidate = os.path.normpath(candidate)
    if not os.path.exists(candidate):
        raise ValueError(
            f"No filename was given, and the last gathered file "
            f"('{last_gathered}') no longer exists. Run "
            f"'python clean.py <path-to-file>' instead."
        )
    return candidate

if __name__ == "__main__":
    try:
        file_path = resolve_input_path(sys.argv[1:], os.getcwd())
    except ValueError as e:
        print(json.dumps({"error": str(e)}))
        sys.exit(1)

    try:
        result = clean_file(file_path)
        if len(sys.argv) == 1:
            result["used_last_gathered_file"] = file_path
        print(json.dumps(result, indent=2))
    except Exception as e:
        print(json.dumps({"error": friendly_error_message(e, file_path)}))
        sys.exit(1)
